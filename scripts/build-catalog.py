# Builds the Cloudbeds data catalog for Rob: every showable data point, flat,
# with a plain-English explanation. Source = live probe of Davenport (44199),
# 2026-06-19. IB-clean formatting (dark header, thin grid, Arial).
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY = "0B1F3A"
NAVY2 = "1F2A44"
LIGHT = "EEF1F6"
RED = "C0392B"
GREY = "5A6473"

thin = Side(style="thin", color="D0D5DD")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, row, ncols, fill=NAVY):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = border

def put_rows(ws, rows, start, ncols, zebra=True):
    r = start
    for row in rows:
        for ci, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=ci, value=val)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = border
        if zebra and (r - start) % 2 == 1:
            for ci in range(1, ncols + 1):
                ws.cell(row=r, column=ci).fill = PatternFill("solid", fgColor=LIGHT)
        r += 1
    return r

wb = Workbook()

# ---------------- Sheet 1: Read Me ----------------
ws = wb.active
ws.title = "Read Me"
ws.sheet_view.showGridLines = False
ws["A1"] = "Stayable — Cloudbeds Dashboard: Available Data Catalog"
ws["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY)
ws["A2"] = "Every data point the dashboard can display, with a plain-English explanation."
ws["A2"].font = Font(name="Arial", size=10, italic=True, color=GREY)

notes = [
    ("Prepared for", "Rob"),
    ("Prepared by", "Kyle Estocapio (BKE)"),
    ("Source", "Live Cloudbeds API probe of Stayable Davenport (Property ID 44199 / API 318197)"),
    ("Date", "2026-06-19"),
    ("Scope", "Same data is available for all 8 properties (each on its own scoped key)."),
    ("How to read", "Sheet 'Available Data' lists every showable metric, one per row. 'Cadence' = "
                    "'Live now' (current moment) or 'Any date range' (pick days/weeks/months)."),
    ("Currently shown", "The dashboard today shows Occupancy % (range) plus the live-snapshot operational "
                        "counts. Everything marked 'Available' can be added on request."),
    ("Privacy note", "Sheet 'Excluded - Guest PII' lists guest-level personal data the API key can technically "
                     "read but that we DO NOT and WILL NOT display — the dashboard is public/view-only. "
                     "See that sheet for the recommended key re-scoping."),
    ("Abbreviations", "ADR = Average Daily Rate; RevPAR = Revenue Per Available Room; "
                      "USALI = Uniform System of Accounts for the Lodging Industry; POS = Point of Sale; "
                      "GL = General Ledger; OOS = Out of Service; PII = Personally Identifiable Information."),
]
r = 4
for k, v in notes:
    ws.cell(row=r, column=1, value=k).font = Font(name="Arial", bold=True, size=10, color=NAVY)
    c = ws.cell(row=r, column=2, value=v)
    c.font = Font(name="Arial", size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 100
for rr in range(4, r):
    ws.row_dimensions[rr].height = 30

# ---------------- Sheet 2: Available Data ----------------
ws2 = wb.create_sheet("Available Data")
ws2.sheet_view.showGridLines = False
headers = ["#", "Data point", "What it is (plain English)", "Category", "Source", "Type", "Cadence", "On dashboard now?"]
for ci, h in enumerate(headers, start=1):
    ws2.cell(row=1, column=h and ci, value=h)
style_header(ws2, 1, len(headers))

LIVE = "Live snapshot (getDashboard)"
OCC = "Data Insights - Occupancy"
RES = "Data Insights - Reservations"
FIN = "Data Insights - Finances"
PAY = "Data Insights - Payments"
INV = "Data Insights - Invoices"
HK = "Data Insights - Housekeeping"
GRP = "Data Insights - Groups & Events"

# (data point, explanation, category, source, type, cadence, on_now)
data = [
    # --- Live operational snapshot ---
    ("Occupancy % (live)", "Share of sellable rooms occupied right now.", "Occupancy", LIVE, "Percent", "Live now", "Yes"),
    ("Rooms Occupied (live)", "Number of rooms occupied right now.", "Occupancy", LIVE, "Count", "Live now", "Yes"),
    ("Sellable Capacity", "Total rooms available to sell (inventory).", "Occupancy", LIVE, "Count", "Live now", "Yes"),
    ("In-House Rooms", "Rooms with a currently checked-in guest.", "Activity", LIVE, "Count", "Live now", "Yes"),
    ("Guests In-House", "Headcount of guests currently in-house.", "Activity", LIVE, "Count", "Live now", "Yes"),
    ("Arrivals (expected today)", "Reservations due to check in today.", "Activity", LIVE, "Count", "Live now", "Yes"),
    ("Arrivals Confirmed", "Arrivals that have actually checked in.", "Activity", LIVE, "Count", "Live now", "Yes"),
    ("Departures (expected today)", "Reservations due to check out today.", "Activity", LIVE, "Count", "Live now", "Yes"),
    ("Departures Confirmed", "Departures that have actually checked out.", "Activity", LIVE, "Count", "Live now", "Yes"),
    ("Stayovers", "Guests staying through (not arriving or leaving today).", "Activity", LIVE, "Count", "Live now", "Yes"),
    ("New Bookings (today)", "Reservations created today.", "Activity", LIVE, "Count", "Live now", "Yes"),
    ("Cancellations (today)", "Reservations cancelled today.", "Activity", LIVE, "Count", "Live now", "Yes"),
    ("Rooms Blocked", "Rooms taken out of availability (blocks + OOS).", "Rooms", LIVE, "Count", "Live now", "Yes"),
    ("% Blocked", "Blocked rooms as a share of capacity.", "Rooms", LIVE, "Percent", "Live now", "Yes"),
    ("Date-Blocked Rooms", "Rooms blocked for date reasons (e.g., holds).", "Rooms", LIVE, "Count", "Live now", "Yes"),
    ("Out-of-Service Rooms", "Rooms unsellable (maintenance/renovation).", "Rooms", LIVE, "Count", "Live now", "Yes"),
    ("Property Local Time", "Current time and timezone at the property.", "Reference", LIVE, "Timestamp", "Live now", "Yes"),

    # --- Occupancy dataset (date-ranged) ---
    ("Occupancy %", "Rooms sold / capacity over the selected dates.", "Occupancy", OCC, "Percent", "Any date range", "Yes"),
    ("Adjusted Occupancy %", "Occupancy after removing blocked/OOS rooms from the denominator.", "Occupancy", OCC, "Percent", "Any date range", "No (available)"),
    ("Rooms Sold", "Total room-nights sold over the range.", "Occupancy", OCC, "Count", "Any date range", "No (available)"),
    ("Rooms Available", "Room-nights available to sell over the range.", "Occupancy", OCC, "Count", "Any date range", "No (available)"),
    ("Capacity", "Total room-nights in inventory over the range.", "Occupancy", OCC, "Count", "Any date range", "No (available)"),
    ("Blocked Rooms", "Room-nights blocked over the range.", "Rooms", OCC, "Count", "Any date range", "No (available)"),
    ("Out-of-Service Rooms", "Room-nights out of service over the range.", "Rooms", OCC, "Count", "Any date range", "No (available)"),
    ("Allotment-Blocked Rooms", "Room-nights held for group/allotment blocks.", "Rooms", OCC, "Count", "Any date range", "No (available)"),
    ("ADR - Average Daily Rate", "Average room revenue per occupied room-night.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("RevPAR - Revenue Per Available Room", "Room revenue per available room-night.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Room Rate", "Base room charge before fees/taxes.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Total Room Revenue", "All room-charge revenue over the range.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Other (Non-Room) Revenue", "Revenue not from room charges (add-ons, items).", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Total Other Revenue", "All non-room revenue, totaled.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Miscellaneous Income", "Income outside standard room/other buckets.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Total Revenue", "Room + other + misc revenue combined.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Room Fees", "Fees charged on room bookings.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Room Taxes", "Taxes collected on room bookings.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Other Room Revenue", "Additional room-related revenue adjustments.", "Revenue", OCC, "Currency", "Any date range", "No (available)"),
    ("Adults", "Adult guest count for the stays in range.", "Guests (aggregate)", OCC, "Count", "Any date range", "No (available)"),
    ("Children", "Children guest count for the stays in range.", "Guests (aggregate)", OCC, "Count", "Any date range", "No (available)"),
    ("Room Guest Count", "Total guests per room over the range.", "Guests (aggregate)", OCC, "Count", "Any date range", "No (available)"),
    ("Market Segment", "Business classification of the demand (e.g., retail, government).", "Segmentation", OCC, "Category", "Any date range", "No (available)"),
    ("Market Group", "High-level demand group (Transient, Group, Contract).", "Segmentation", OCC, "Category", "Any date range", "No (available)"),
    ("Room Type", "Room product sold (e.g., King, Double Queen).", "Segmentation", OCC, "Category", "Any date range", "No (available)"),
    ("Room Type Category", "Private vs. shared room classification.", "Segmentation", OCC, "Category", "Any date range", "No (available)"),

    # --- Reservations / booking pace ---
    ("Room Nights", "Total room-nights booked across reservations.", "Reservations", RES, "Count", "Any date range", "No (available)"),
    ("Room Count", "Number of rooms per reservation.", "Reservations", RES, "Count", "Any date range", "No (available)"),
    ("Booking Window (lead time)", "Days between booking and check-in (pace indicator).", "Reservations", RES, "Number", "Any date range", "No (available)"),
    ("Reservation Grand Total", "Full value of a reservation.", "Reservations", RES, "Currency", "Any date range", "No (available)"),
    ("Reservation Paid Amount", "Amount already paid on reservations.", "Reservations", RES, "Currency", "Any date range", "No (available)"),
    ("Reservation Balance Due", "Outstanding amount owed on reservations.", "Reservations", RES, "Currency", "Any date range", "No (available)"),
    ("Cancellation Fee", "Fees charged on cancellations.", "Reservations", RES, "Currency", "Any date range", "No (available)"),
    ("No-Show Fee", "Fees charged on no-shows.", "Reservations", RES, "Currency", "Any date range", "No (available)"),
    ("Channel Commission Amount", "Commission owed to booking channels/OTAs.", "Reservations", RES, "Currency", "Any date range", "No (available)"),
    ("Commission Percent", "Commission rate applied by the channel.", "Reservations", RES, "Percent", "Any date range", "No (available)"),
    ("Reservation Source (channel)", "Where the booking came from (Phone, Expedia, etc.).", "Channel mix", RES, "Category", "Any date range", "No (available)"),
    ("Reservation Source Category", "Channel grouping (Direct, OTA, Travel Agent, etc.).", "Channel mix", RES, "Category", "Any date range", "No (available)"),
    ("Reservation Status", "Confirmed / In-House / Checked-Out / Cancelled / No-Show.", "Reservations", RES, "Category", "Any date range", "No (available)"),
    ("Rate Plan (public name)", "The customer-facing rate plan booked.", "Reservations", RES, "Category", "Any date range", "No (available)"),
    ("Meal Plan Included", "Whether a meal plan is attached.", "Reservations", RES, "Flag", "Any date range", "No (available)"),
    ("Total Reservation Taxes", "Taxes across reservations.", "Reservations", RES, "Currency", "Any date range", "No (available)"),
    ("Total Reservation Fees", "Fees across reservations.", "Reservations", RES, "Currency", "Any date range", "No (available)"),

    # --- Finances / revenue detail ---
    ("Debits (charges)", "Amounts charged to folios.", "Finance", FIN, "Currency", "Any date range", "No (available)"),
    ("Credits (payments/credits)", "Payments and credits applied to folios.", "Finance", FIN, "Currency", "Any date range", "No (available)"),
    ("Transaction Amount (net)", "Net of charges minus payments/credits.", "Finance", FIN, "Currency", "Any date range", "No (available)"),
    ("Benchmarking Transaction Type", "USALI revenue category for benchmarking.", "Finance", FIN, "Category", "Any date range", "No (available)"),
    ("Transaction Type", "Nature of the transaction (charge, payment, tax, fee...).", "Finance", FIN, "Category", "Any date range", "No (available)"),
    ("Fee Type", "Type of fee charged.", "Finance", FIN, "Category", "Any date range", "No (available)"),
    ("Tax Type", "Type of tax collected.", "Finance", FIN, "Category", "Any date range", "No (available)"),
    ("Add-on Item", "Extras sold with a reservation (e.g., parking).", "Finance", FIN, "Category", "Any date range", "No (available)"),
    ("Item & Service Category", "Grouping of sold items/services.", "Finance", FIN, "Category", "Any date range", "No (available)"),
    ("POS Charge Category", "Point-of-Sale charge grouping.", "Finance", FIN, "Category", "Any date range", "No (available)"),
    ("Payment Method", "How guests paid (card, cash, etc.).", "Finance", FIN, "Category", "Any date range", "No (available)"),
    ("Payment Type", "Classification of the payment.", "Finance", FIN, "Category", "Any date range", "No (available)"),
    ("General Ledger Name", "Custom GL account mapping (if configured).", "Finance", FIN, "Category", "Any date range", "No (available)"),

    # --- Payments / payouts ---
    ("Captured Amount", "Payment amount successfully captured.", "Payments", PAY, "Currency", "Any date range", "No (available)"),
    ("Refunded Amount", "Amount refunded to guests.", "Payments", PAY, "Currency", "Any date range", "No (available)"),
    ("Payment Net", "Payment net of processing fees.", "Payments", PAY, "Currency", "Any date range", "No (available)"),
    ("Payment Fee", "Processing fee on a payment.", "Payments", PAY, "Currency", "Any date range", "No (available)"),
    ("Payout Amount", "Amount paid out to the property.", "Payments", PAY, "Currency", "Any date range", "No (available)"),
    ("Payout Status", "Status of a payout (paid, pending, etc.).", "Payments", PAY, "Category", "Any date range", "No (available)"),
    ("Payment Gateway Status", "Result from the payment gateway.", "Payments", PAY, "Category", "Any date range", "No (available)"),

    # --- Invoices ---
    ("Invoice Grand Total", "Total value of an issued invoice.", "Invoices", INV, "Currency", "Any date range", "No (available)"),
    ("Invoice Total Tax", "Tax portion of an invoice.", "Invoices", INV, "Currency", "Any date range", "No (available)"),
    ("Invoice Balance Due", "Amount still owed on an invoice.", "Invoices", INV, "Currency", "Any date range", "No (available)"),
    ("Invoice Status", "Open / paid / void, etc.", "Invoices", INV, "Category", "Any date range", "No (available)"),
    ("Invoice Age", "Days since invoice / due date (aging).", "Invoices", INV, "Number", "Any date range", "No (available)"),
    ("Reservation Room Nights (invoice)", "Room nights tied to the invoice.", "Invoices", INV, "Count", "Any date range", "No (available)"),

    # --- Housekeeping & maintenance ---
    ("Room Condition", "Clean / dirty / inspected status of a room.", "Housekeeping", HK, "Category", "Live / by date", "No (available)"),
    ("Room Status", "Occupancy/housekeeping status of a room.", "Housekeeping", HK, "Category", "Live / by date", "No (available)"),
    ("Frontdesk Status", "Front-desk view of room state.", "Housekeeping", HK, "Category", "Live / by date", "No (available)"),
    ("Do Not Disturb", "DND flag on a room.", "Housekeeping", HK, "Flag", "Live / by date", "No (available)"),
    ("Housekeeper Assigned", "Staff member assigned to a room.", "Housekeeping", HK, "Text", "Live / by date", "No (available)"),
    ("Room Count by Status", "Number of rooms in each status.", "Housekeeping", HK, "Count", "Live / by date", "No (available)"),

    # --- Groups & events ---
    ("Allotment Block Count", "Rooms held in a group block.", "Groups", GRP, "Count", "Any date range", "No (available)"),
    ("Remaining Block Count", "Unsold rooms left in a block.", "Groups", GRP, "Count", "Any date range", "No (available)"),
    ("Rooms Sold (group)", "Rooms picked up from a group block.", "Groups", GRP, "Count", "Any date range", "No (available)"),
    ("Blocked Room Revenue", "Revenue tied to blocked group rooms.", "Groups", GRP, "Currency", "Any date range", "No (available)"),
    ("Group Total Room Revenue", "Total room revenue from a group/event.", "Groups", GRP, "Currency", "Any date range", "No (available)"),
    ("Allotment Block Status", "Status of the group block.", "Groups", GRP, "Category", "Any date range", "No (available)"),
    ("Event Name", "Name of the group/event.", "Groups", GRP, "Text", "Any date range", "No (available)"),
    ("Event Status", "Status of the event.", "Groups", GRP, "Category", "Any date range", "No (available)"),
]

rows = [[i + 1, *d] for i, d in enumerate(data)]
end = put_rows(ws2, rows, 2, len(headers))
ws2.freeze_panes = "A2"
widths = [4, 30, 52, 18, 28, 12, 16, 16]
for i, w in enumerate(widths, start=1):
    ws2.column_dimensions[chr(64 + i)].width = w

# ---------------- Sheet 3: Excluded - Guest PII ----------------
ws3 = wb.create_sheet("Excluded - Guest PII")
ws3.sheet_view.showGridLines = False
ws3["A1"] = "Excluded from the dashboard — Guest-level personal data (PII)"
ws3["A1"].font = Font(name="Arial", bold=True, size=12, color=RED)
ws3["A2"] = ("These fields are technically readable by the current API key (verified live: getGuestList "
             "returned 100 guest records). They are NEVER displayed on the dashboard — it is public/view-only. "
             "Recommend re-scoping the keys to remove Guest + Data Insights Guests access so this is impossible "
             "by design, not just by policy.")
ws3["A2"].font = Font(name="Arial", size=10, color=GREY)
ws3["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws3.merge_cells("A2:C2")
ws3.row_dimensions[2].height = 56

ph = ["Guest data field", "What it is", "Where it's reachable"]
for ci, h in enumerate(ph, start=1):
    ws3.cell(row=4, column=ci, value=h)
style_header(ws3, 4, len(ph), fill=NAVY2)

V13G = "v1.3 getGuestList / DI Guests, Reservations"
pii = [
    ("Guest Full Name / First / Surname", "Identifies the individual guest.", V13G),
    ("Email", "Guest email address.", V13G),
    ("Phone / Mobile", "Guest contact numbers.", V13G),
    ("Date of Birth", "Guest DOB.", V13G),
    ("Home Address / City / State / ZIP / Country", "Guest residential address.", V13G),
    ("Document Number / Type / Expiration", "Passport / ID document details.", V13G),
    ("Gender", "Guest gender.", V13G),
    ("Marketing Opt-In", "Whether the guest opted into marketing.", V13G),
    ("Tax ID (guest/company)", "Guest or company tax identifier.", "DI Reservations / Invoices"),
    ("Card Type / Card Last 4", "Partial payment card detail.", "DI Finances / Payments / Reservations"),
]
prows = [[a, b, c] for (a, b, c) in pii]
put_rows(ws3, prows, 5, len(ph))
ws3.column_dimensions["A"].width = 40
ws3.column_dimensions["B"].width = 44
ws3.column_dimensions["C"].width = 38

out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "outputs")
os.makedirs(out_dir, exist_ok=True)
path = os.path.join(out_dir, "CloudbedsDataCatalog_Stayable_061926.xlsx")
wb.save(path)
print("saved:", path)
print("metrics:", len(data), "| pii rows:", len(pii))
