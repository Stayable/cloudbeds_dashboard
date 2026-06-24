# Light edit (no rebuild): add an "Approve (Rob)" column to the Available Data
# sheet and flag the previously-removed revenue metrics as re-add candidates.
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "0B1F3A"
GOLD = "FFF3CD"
thin = Side(style="thin", color="D0D5DD")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    "outputs", "CloudbedsDataCatalog_Stayable_061926.xlsx")
wb = load_workbook(path)
ws = wb["Available Data"]

approve_col = ws.max_column + 1  # col 9 (I)
last_row = ws.max_row

# header
h = ws.cell(row=1, column=approve_col, value="Approve (Rob)")
h.fill = PatternFill("solid", fgColor=NAVY)
h.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
h.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
h.border = border

# blank approval cells + Yes/No dropdown
for r in range(2, last_row + 1):
    c = ws.cell(row=r, column=approve_col, value="")
    c.font = Font(name="Arial", size=10)
    c.alignment = Alignment(horizontal="center", vertical="top")
    c.border = border
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(dv)
col_letter = ws.cell(row=1, column=approve_col).column_letter
dv.add(f"{col_letter}2:{col_letter}{last_row}")
ws.column_dimensions[col_letter].width = 14

# flag previously-removed revenue metrics as re-add candidates (col 8 = "On dashboard now?")
readd = {"ADR", "RevPAR", "Total Room Revenue", "Total Revenue"}
for r in range(2, last_row + 1):
    name = str(ws.cell(row=r, column=2).value or "")
    if any(name.startswith(x) for x in readd):
        ws.cell(row=r, column=8, value="No - removed earlier (re-add candidate)")
        for ci in range(1, approve_col + 1):
            ws.cell(row=r, column=ci).fill = PatternFill("solid", fgColor=GOLD)

# add a short note to Read Me about the approval column + revenue feasibility
rm = wb["Read Me"]
nr = rm.max_row + 2
for label, text in [
    ("Approval", "Available Data tab now has an 'Approve (Rob)' column (Yes/No). Rob marks which "
                 "data points to add to the dashboard."),
    ("Revenue (re-add)", "ADR, RevPAR, Total Room Revenue and Total Revenue were originally removed from the "
                         "dashboard UI. They are highlighted as re-add candidates. ADR/RevPAR aggregate "
                         "natively; exact room revenue is obtained by summing the daily Data Insights values."),
]:
    rm.cell(row=nr, column=1, value=label).font = Font(name="Arial", bold=True, size=10, color=NAVY)
    cc = rm.cell(row=nr, column=2, value=text)
    cc.font = Font(name="Arial", size=10)
    cc.alignment = Alignment(wrap_text=True, vertical="top")
    rm.row_dimensions[nr].height = 30
    nr += 1

wb.save(path)
print("updated:", path, "| approve col:", col_letter, "| rows:", last_row)
