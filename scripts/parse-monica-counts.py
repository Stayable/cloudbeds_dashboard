# Parse daily occupancy COUNTS from Monica's "Occ% Report" workbook for a given
# year, into JSON for backfill-counts-from-monica.mjs to load into Neon.
#
# WHY this exists: historical daily counts can't be faithfully reconstructed from
# Cloudbeds (dataset-3 rate-plan/status is current-state → retroactive
# reclassification drifts the transient/lease split and daily totals; verified
# 07/25/26). Monica's workbook holds the FROZEN daily actuals, so we source the
# count cells from it. Revenue is NOT sourced here — it's backfilled separately
# and exactly from Cloudbeds dataset-1 (run-backfill.mjs).
#
# CRITICAL layout note: each per-property tab stacks yearly blocks. A given year
# appears TWICE — once as the "current year" column (idx1) in its own block, and
# once as the "last year" column (idx2) of the NEXT year's block. Only the
# LAST-YEAR-column copy (Block C: date in idx2, idx1 empty/next-year) is the
# SETTLED/FINAL actual — it ties to the summary tab exactly. The current-year
# copy (Block B, idx1) is the as-captured/on-the-books version and is incomplete.
# So we read Block C.
#
# Usage: python scripts/parse-monica-counts.py <workbook.xlsx> <year> <cap_YYYY-MM-DD> [out.json]
#   cap = last day to include (exclusive of cron-banked days; e.g. day before the
#   daily cron started banking counts). Example:
#   python scripts/parse-monica-counts.py "Occupancy Report/Occ% Report as of July 21.xlsx" 2026 2026-07-20

import sys, json, datetime
import openpyxl
from openpyxl.utils.datetime import from_excel

WB = sys.argv[1]
YEAR = int(sys.argv[2])
CAP = datetime.date.fromisoformat(sys.argv[3])
OUT = sys.argv[4] if len(sys.argv) > 4 else "monica-counts.json"

TABS = {
    "4645 - Lakeland": "LL", "6802 - JS West": "JW", "2295 - KS East": "KE",
    "5399 - KS West": "KW", "8700 - Orlando": "OR", "2535 - St. Augustine": "SA",
    "44199 - Davenport": "DP", "812 - JS North": "JN",
}
NAME = {"Lakeland": "LL", "Jacksonville West": "JW", "Kissimmee East": "KE",
        "Kissimmee West": "KW", "Orlando": "OR", "St. Augustine": "SA",
        "Davenport": "DP", "Jacksonville North": "JN"}
# Column indices (0-based) within a daily row — shared header template:
#   idx9 Room Inventory · idx11 OOO · idx13 Transient-CB (nights) ·
#   idx17 Lease Total (nights) · idx18 Other Blocks · idx19 Occupied Rooms
C_INV, C_OOO, C_TRANS, C_LEASE, C_OTHER, C_OCC = 9, 11, 13, 17, 18, 19


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, (int, float)) and 40000 < v < 60000:
        try:
            return from_excel(v).date()
        except Exception:
            return None
    return None


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return 0.0


wb = openpyxl.load_workbook(WB, data_only=True, read_only=True)

# Summary-tab YTD occupied/transient/lease for validation.
summ = {}
cur = None
for r in wb["Occ% Report"].iter_rows(values_only=True):
    b = r[1] if len(r) > 1 else None
    if isinstance(b, str) and b.startswith("Stayable "):
        cur = b.replace("Stayable ", "").strip(); summ[cur] = {}
    elif cur and b in ("Occupied", "Transient", "Lease") and b not in summ[cur]:
        summ[cur][b] = r[26] if len(r) > 26 else None
sumC = {NAME[k]: v for k, v in summ.items() if k in NAME}

out = []
val = {}
for tab, code in TABS.items():
    ws = wb[tab]
    seen = set(); occ = tr = le = 0
    for r in ws.iter_rows(values_only=True):
        d2 = as_date(r[2] if len(r) > 2 else None)
        d1v = r[1] if len(r) > 1 else None
        if not d2 or d2.year != YEAR or d2 > CAP:
            continue
        # Block C only: date is in the last-year column (idx2); idx1 empty or next year.
        if not (d1v is None or (as_date(d1v) and as_date(d1v).year == YEAR + 1)):
            continue
        k = d2.isoformat()
        if k in seen:
            continue
        seen.add(k)
        t = int(round(num(r[C_TRANS]))); l = int(round(num(r[C_LEASE])))
        out.append({
            "code": code, "date": k,
            "transientNights": t, "leaseNights": l,
            "otherBlocks": int(round(num(r[C_OTHER]))),
            "ooo": int(round(num(r[C_OOO]))),
            "inventory": int(round(num(r[C_INV]))),
        })
        occ += int(round(num(r[C_OCC]))); tr += t; le += l
    val[code] = {"days": len(seen), "occ": occ, "tr": tr, "le": le}

with open(OUT, "w") as f:
    json.dump(out, f)

print(f"Parsed {len(out)} property-days (Block C, {YEAR} through {CAP}) -> {OUT}\n")
print(f"{'code':<5}{'days':>5}{'occ(parsed)':>13}{'occ(summary)':>14}  (summary is thru report as-of; expect ~1 day more)")
for code, v in val.items():
    s = sumC.get(code, {})
    print(f"{code:<5}{v['days']:>5}{v['occ']:>13,}{(s.get('Occupied') or 0):>14,.0f}")
